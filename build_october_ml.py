"""
build_october_ml.py
-------------------
Splice October onto the PUBLISHED ML Forecast book. Same rule as the other two:
history is immutable, only the tail is rewritten.

Why ML needs this more than the others. The ML pipeline is a walk-forward
LightGBM feeding a SOM/KMeans clustering step and then the EGP optimiser, and it
is not deterministic: re-running it today moved the last nine months of history
by up to 5.4pp (2026-07 went -1.80% -> +3.56%) and produced a September book
holding SUNTV at 2164 shares, which was never held. So nothing is taken from the
October run except the October selection itself.

  * monthly_detail / Summary FULL for trade months <= 2026-08 : untouched.
  * Trade 2026-09 (September) : published as a 0.00% stub. Realised here from
    the PUBLISHED September book (the 12 names in ml_data.js current_portfolio)
    priced over September.
  * Trade 2026-10 (October) : the October run's 7-name selection, with Prev Qty
    re-pointed at the published September book and an exit row for each of the
    10 September names October drops.

One honest limitation. som-style monthly P&L includes a gap term for names held
in the PREVIOUS month but not the current one -- (trade-month open minus prior
close) on the outgoing quantity. Reconstructing it needs the AUGUST book, and
that is not recoverable: ml_data.js carries only the current book, its
exec_history is empty, and the pipeline's own workbooks were overwritten by
today's re-runs. Every published September holding records prev_qty = 0, so the
gap term is taken as zero. September's realised figure is therefore the book's
own September performance net of transaction cost, and is reported as such.

    python build_october_ml.py            # -> scratch/spliced/
    python build_october_ml.py --install  # also copy over the live workbooks
"""
import json
import os
import re
import shutil
import sys

import numpy as np
import openpyxl
import pandas as pd

MAIN = os.path.dirname(os.path.abspath(__file__))
ML_DIR = os.environ.get("ML_PIPELINE_DIR", r"d:\PC2546\portfolio")
OCT = os.path.join(MAIN, "scratch", "oct_run")
OUT = os.path.join(MAIN, "scratch", "spliced")

SUMMARY = "Sharpe_ML_Forecast_NIFTY500_Summary.xlsx"
CURRENT = "Current_Portfolio_ML_Forecast_NIFTY500.xlsx"
OCT_SUMMARY = os.path.join(OCT, "ml_summary_oct.xlsx")
OCT_CURRENT = os.path.join(OCT, "ml_current_oct.xlsx")

PRICES = os.path.join(ML_DIR, "NIFTY500")
BENCH = os.path.join(ML_DIR, "NSE_CNX500, 1D.csv")

CAP = 10_000_000
COST_PER_TRADE = 0.002

SIG_REALIZE = "2026-08"   # traded September
SIG_NEW = "2026-09"       # traded October
MONTH_RE = re.compile(r"^\d{4}-\d{2}$")

_norm = lambda x: re.sub(r"\s+", " ", str(x)).strip()
_bars = {}


def monthly(sym):
    """month Period -> [open, close]."""
    if sym in _bars:
        return _bars[sym]
    p = os.path.join(PRICES, sym + "_1d_max.csv")
    if not os.path.exists(p):
        _bars[sym] = {}
        return {}
    d = pd.read_csv(p)
    dc = "Date" if "Date" in d.columns else d.columns[0]
    d[dc] = pd.to_datetime(d[dc], format="%d-%m-%Y", errors="coerce")
    d = d.dropna(subset=[dc]).sort_values(dc)
    out = {}
    for _, r in d.iterrows():
        k = pd.Period(r[dc], "M")
        try:
            o, c = float(r["Open"]), float(r["Close"])
        except (TypeError, ValueError, KeyError):
            continue
        if c != c:
            continue
        if k not in out:
            out[k] = [o, c]
        out[k][1] = c
    _bars[sym] = out
    return out


def num(v):
    try:
        f = float(v)
        return 0.0 if f != f else f
    except (TypeError, ValueError):
        return 0.0


def published_book():
    """The September book actually published: {symbol: (qty, prev_qty)}."""
    for ln in open(os.path.join(MAIN, "ml_data.js"), encoding="utf-8"):
        if "DASHBOARD_DATA.ml_forecast" in ln and "=" in ln:
            obj = json.loads(ln[ln.index("=") + 1:].strip().rstrip(";"))
            break
    else:
        sys.exit("[error] ml_data.js has no DASHBOARD_DATA.ml_forecast line")
    return ({str(r["symbol"]).strip(): (int(r.get("qty") or 0),
                                        int(r.get("prev_qty") or 0))
             for r in obj["current_portfolio"]}, obj)


def realised_return(book, port_month, trade_month):
    """som-style: [gap + current P&L - txn] / CAP, over the published book.

    gap uses the recorded prev_qty (all zero in the published book -- see the
    module docstring), current uses the held quantity.
    """
    pm, tm = pd.Period(port_month, "M"), pd.Period(trade_month, "M")
    pnl = txn = 0.0
    missing = []
    for s, (q, pq) in book.items():
        md = monthly(s)
        hist = [k for k in md if k <= pm] if md else []
        if not hist or tm not in md:
            missing.append(s)
            continue
        prev_close = md[max(hist)][1]
        buy_px = md[tm][0] or prev_close
        if buy_px <= 0:
            buy_px = prev_close
        sell_px = md[tm][1]
        pnl += (buy_px - prev_close) * pq + (sell_px - buy_px) * q
        txn += abs(q - pq) * buy_px
    return (pnl - txn * COST_PER_TRADE) / CAP, missing


def bench_month(trade_month):
    d = pd.read_csv(BENCH)
    dc = "time" if "time" in d.columns else "Date"
    cc = "close" if "close" in d.columns else "Close"
    d[dc] = pd.to_datetime(d[dc], errors="coerce")
    d = d.dropna(subset=[dc, cc]).set_index(dc).sort_index()
    m = d[cc].resample("ME").last().pct_change().dropna()
    m.index = m.index.to_period("M")
    tm = pd.Period(trade_month, "M")
    return float(m.loc[tm]) if tm in m.index else 0.0


def hdr_of(ws):
    hr = next(r for r in range(1, 30)
              if any(_norm(ws.cell(r, c).value) == "Portfolio Month"
                     for c in range(1, ws.max_column + 1)))
    ci = {}
    for c in range(1, ws.max_column + 1):
        v = _norm(ws.cell(hr, c).value)
        if v and v != "nan":
            ci[v] = c
    return hr, ci


def month_rows(ws, hr, ci):
    col = ci["Portfolio Month"]
    return {str(ws.cell(r, col).value).strip(): r
            for r in range(hr + 1, ws.max_row + 1)
            if ws.cell(r, col).value
            and MONTH_RE.match(str(ws.cell(r, col).value).strip())}


def selection(path):
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r or not r[0]:
            continue
        out[str(r[0]).strip()] = {"tgt": int(num(r[3])), "prev": int(num(r[4])),
                                  "w": num(r[5]), "px": num(r[6])}
    return ws.title, out


def main():
    install = "--install" in sys.argv
    os.makedirs(OUT, exist_ok=True)
    for p in (os.path.join(ML_DIR, SUMMARY), os.path.join(ML_DIR, CURRENT),
              OCT_SUMMARY, OCT_CURRENT):
        if not os.path.exists(p):
            sys.exit("[error] missing " + p)

    book, _ = published_book()
    trade_realize = (pd.Period(SIG_REALIZE, "M") + 1).strftime("%Y-%m")
    trade_new = (pd.Period(SIG_NEW, "M") + 1).strftime("%Y-%m")

    print("=== ML Forecast ===")
    ret, miss = realised_return(book, SIG_REALIZE, trade_realize)
    ben = bench_month(trade_realize)
    if miss:
        print("    [warn] no September bar for " + ", ".join(miss))

    # ---------------- Summary FULL (and the windowed copies) ----------------
    dst = os.path.join(OUT, SUMMARY)
    shutil.copy2(os.path.join(ML_DIR, SUMMARY), dst)
    wbo = openpyxl.load_workbook(OCT_SUMMARY, data_only=True)
    wb = openpyxl.load_workbook(dst)

    for sheet in [s for s in wb.sheetnames if s.startswith("Summary")]:
        ws = wb[sheet]
        try:
            hr, ci = hdr_of(ws)
        except StopIteration:
            continue
        mr = month_rows(ws, hr, ci)
        if SIG_REALIZE not in mr:
            continue
        prev_sig = sorted(m for m in mr if m < SIG_REALIZE)[-1]
        carry = {k: num(ws.cell(mr[prev_sig], ci[k]).value)
                 for k in ("Port Cumul %", "Bench Cumul %",
                           "Compounded Capital (Rs.)") if k in ci}
        r = mr[SIG_REALIZE]
        old = num(ws.cell(r, ci["Port Return %"]).value)
        ws.cell(r, ci["Port Return %"], ret)
        ws.cell(r, ci["Bench Return %"], ben)
        if "Net PnL (Rs.)" in ci:
            ws.cell(r, ci["Net PnL (Rs.)"], ret * CAP)
        if "Alpha %" in ci:
            ws.cell(r, ci["Alpha %"], ret - ben)
        if "Beat Bench?" in ci:
            ws.cell(r, ci["Beat Bench?"], "YES" if ret > ben else "NO")
        pc = (1 + carry.get("Port Cumul %", 0.0)) * (1 + ret) - 1
        bc = (1 + carry.get("Bench Cumul %", 0.0)) * (1 + ben) - 1
        cap = carry.get("Compounded Capital (Rs.)", CAP) * (1 + ret)
        for k, v in (("Port Cumul %", pc), ("Bench Cumul %", bc),
                     ("Compounded Capital (Rs.)", cap)):
            if k in ci:
                ws.cell(r, ci[k], v)

        # append the October row from the October run's own summary
        wso = wbo[sheet] if sheet in wbo.sheetnames else wbo["Summary FULL"]
        hro, cio = hdr_of(wso)
        mro = month_rows(wso, hro, cio)
        if SIG_NEW not in mro:
            continue
        ins = max(mr.values()) + 1
        ws.insert_rows(ins)
        for k, c in ci.items():
            if k in cio:
                ws.cell(ins, c, wso.cell(mro[SIG_NEW], cio[k]).value)
        ws.cell(ins, ci["Portfolio Month"], SIG_NEW)
        if "Trade Month" in ci:
            ws.cell(ins, ci["Trade Month"], trade_new)
        for k in ("Port Return %", "Bench Return %", "Alpha %", "Net PnL (Rs.)"):
            if k in ci:
                ws.cell(ins, ci[k], 0)
        for k, v in (("Port Cumul %", pc), ("Bench Cumul %", bc),
                     ("Compounded Capital (Rs.)", cap)):
            if k in ci:
                ws.cell(ins, ci[k], v)
        if "Beat Bench?" in ci:
            ws.cell(ins, ci["Beat Bench?"], "")
        if sheet == "Summary FULL":
            print("    Sept (trade %s) realised: %+.3f%% (was a %+.3f%% stub)"
                  "   bench %+.3f%%" % (trade_realize, ret * 100, old * 100,
                                        ben * 100))
    wb.save(dst)
    print("    -> " + dst)

    # ---------------- Current selection ----------------
    sep_name, _sep = selection(os.path.join(ML_DIR, CURRENT))
    oct_name, octb = selection(OCT_CURRENT)
    exits = sorted(s for s in book if s not in octb)

    cdst = os.path.join(OUT, CURRENT)
    shutil.copy2(OCT_CURRENT, cdst)
    wbc = openpyxl.load_workbook(cdst)
    wsc = wbc[wbc.sheetnames[0]]
    fixed = 0
    r = 4
    while r <= wsc.max_row and wsc.cell(r, 1).value:
        sym = str(wsc.cell(r, 1).value).strip()
        tgt = int(num(wsc.cell(r, 4).value))
        real_prev = book.get(sym, (0, 0))[0]
        if int(num(wsc.cell(r, 5).value)) != real_prev:
            fixed += 1
        wsc.cell(r, 5, real_prev)
        d = tgt - real_prev
        wsc.cell(r, 3, abs(d))
        wsc.cell(r, 2, "HOLD" if d == 0 else ("BUY" if d > 0 else "SELL"))
        r += 1
    for sym in exits:
        q = book[sym][0]
        md = monthly(sym)
        px = md[max(md)][1] if md else 0.0
        for c, v in ((1, sym), (2, "EXIT SELL"), (3, q), (4, 0), (5, q),
                     (6, 0.0), (7, px)):
            wsc.cell(r, c, v)
        r += 1
    wbc.save(cdst)
    print("    Oct book: %d holdings, Prev Qty corrected on %d row(s), "
          "%d exit row(s)" % (len(octb), fixed, len(exits)))
    print("      book out: " + ", ".join("%s x%d" % (s, book[s][0])
                                         for s in exits))
    print("    -> " + cdst)

    if install:
        for f in (SUMMARY, CURRENT):
            shutil.copy2(os.path.join(OUT, f), os.path.join(ML_DIR, f))
            print("    installed -> " + os.path.join(ML_DIR, f))
    else:
        print("\nnot installed; rerun with --install")


if __name__ == "__main__":
    main()
