"""
reprice_october_open.py
-----------------------
Re-size the October books at today's 09:15 opening price, keeping the selection
and the weights that the 30-09-2026 signal produced.

The split matters. Which stocks are held and at what weight is a 30-September
decision -- the EGP optimiser ran on data through the 30th and nothing here
touches it. How many SHARES that weight buys is a 1-October execution fact, and
the engine's own rule is

    qty = floor(INITIAL_CAPITAL * weight / trade_month_open)

(som_hedge.py line 529). When the books were spliced on 30-09 there was no
October bar yet, so buy_px fell back to the September close and every quantity
was sized off a stale price. This re-sizes them off the real open.

Deliberately does NOT write today's bar into the price CSVs. update_stocks.py
appends by date and skips a date it already has, so a partial 09:15 bar written
now would still be sitting there tonight as the permanent 01-10 bar, with the
day's real high/low/close lost. The opening price is fetched live for sizing
only; tonight's job appends the settled bar as usual.

    python reprice_october_open.py            # -> rewrites scratch/spliced/
    python reprice_october_open.py --dry-run
"""
import csv
import datetime as _dt
import os
import sys
from datetime import datetime

import numpy as np
import openpyxl
import pandas as pd
import yfinance as yf

MAIN = os.path.dirname(os.path.abspath(__file__))
SPL = os.path.join(MAIN, "scratch", "spliced")

# Where the settled daily bars live. The September CLOSE comes from here, never
# from the book's implied sizing price (weight * capital / qty): that implied
# price is the price the position was SIZED at, i.e. the previous month's close,
# so using it turned the one-day overnight gap into a whole month's move and
# showed the October row opening at -6.8%.
PRICE_FOLDERS = ["nifty50_host", "nifty500_host", "TOTAL_STOCKS",
                 "nifty_50_october", "nifty_500_october", "TOTAL_STOCKS_October",
                 "hq_quarterly_universe", "High_Quality_October",
                 r"d:\PC2546\portfolio\NIFTY500"]
_closes = {}

CAP = 10_000_000
COST_PER_TRADE = 0.002
SIG_NEW = "2026-09"          # signal month; the book is traded in October
TRADE = "2026-10"

HEDGE = ["Hedge_nifty50.xlsx", "Hedge_nifty500.xlsx", "Hedge_Pro_Summary_759.xlsx"]
CURRENT_BOOKS = ["SOM_HQ_Quarterly_v2_Current.xlsx",
                 "Current_Portfolio_ML_Forecast_NIFTY500.xlsx"]

DRY = "--dry-run" in sys.argv

# A few names do not resolve on .NS under the symbol the books use.
SUFFIX_OVERRIDE = {"515008": "515008.BO", "543971": "543971.BO",
                   "543619": "543619.BO"}


def num(v):
    try:
        f = float(v)
        return 0.0 if f != f else f
    except (TypeError, ValueError):
        return 0.0


def clean(s):
    return str(s).replace("_1d_max", "").strip()


def sep_close(sym):
    """The settled 30-09-2026 close from the local CSVs, or None."""
    if sym in _closes:
        return _closes[sym]
    val = None
    for fol in PRICE_FOLDERS:
        p = os.path.join(MAIN, fol, sym + "_1d_max.csv") \
            if not os.path.isabs(fol) else os.path.join(fol, sym + "_1d_max.csv")
        if not os.path.exists(p):
            continue
        rows = list(csv.reader(open(p, encoding="utf-8", errors="replace")))
        if not rows or "Date" not in rows[0] or "Close" not in rows[0]:
            continue
        di, ci = rows[0].index("Date"), rows[0].index("Close")
        best = None
        for x in rows[1:]:
            if len(x) <= max(di, ci):
                continue
            try:
                d = _dt.datetime.strptime(x[di].strip(), "%d-%m-%Y")
                c = float(x[ci])
            except ValueError:
                continue
            if d.year == 2026 and d.month == 9 and (best is None or d > best[0]):
                best = (d, c)
        if best:
            val = best[1]
            break
    _closes[sym] = val
    return val


def fetch_opens(symbols):
    """{symbol: (open, last)} for today. Falls back to .BO, then gives up."""
    out = {}
    todo = sorted(symbols)
    for attempt, suffix in ((1, ".NS"), (2, ".BO")):
        if not todo:
            break
        tickers = {s: SUFFIX_OVERRIDE.get(s, s + suffix) for s in todo}
        try:
            df = yf.download(list(tickers.values()), period="2d", interval="1d",
                             auto_adjust=False, progress=False, group_by="ticker",
                             threads=True)
        except Exception as e:                       # noqa: BLE001
            print("    [warn] batch %d failed (%s); falling back per symbol"
                  % (attempt, type(e).__name__))
            df = None
        still = []
        for s, tk in tickers.items():
            o = c = None
            try:
                sub = df[tk] if df is not None and tk in df.columns.get_level_values(0) else None
                if sub is not None:
                    sub = sub.dropna(subset=["Open"])
                    if not sub.empty:
                        last_idx = sub.index[-1]
                        if last_idx.strftime("%Y-%m") == TRADE:
                            o, c = float(sub["Open"].iloc[-1]), float(sub["Close"].iloc[-1])
            except Exception:                        # noqa: BLE001
                o = c = None
            if o is None:
                try:
                    h = yf.Ticker(tk).history(period="2d", interval="1d",
                                             auto_adjust=False)
                    h = h.dropna(subset=["Open"])
                    if not h.empty and h.index[-1].strftime("%Y-%m") == TRADE:
                        o, c = float(h["Open"].iloc[-1]), float(h["Close"].iloc[-1])
                except Exception:                    # noqa: BLE001
                    pass
            if o and o > 0:
                out[s] = (o, c if c and c > 0 else o)
            else:
                still.append(s)
        todo = still
    return out, todo


def table(ws, hdr_label="Stock"):
    """Locate the header row of a Port_ sheet and map column name -> index."""
    hr = next(r for r in range(1, 12)
              if str(ws.cell(r, 1).value).strip() == hdr_label)
    ci = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(hr, c).value
        if v:
            ci[str(v).strip()] = c
    return hr, ci


def main():
    # ---- 1. collect every symbol that needs an opening price ----
    want = set()
    for f in HEDGE:
        wb = openpyxl.load_workbook(os.path.join(SPL, f), data_only=True)
        ws = wb["Port_" + SIG_NEW]
        hr, ci = table(ws)
        for r in range(hr + 1, ws.max_row + 1):
            if ws.cell(r, ci["Stock"]).value:
                want.add(clean(ws.cell(r, ci["Stock"]).value))
    for f in CURRENT_BOOKS:
        wb = openpyxl.load_workbook(os.path.join(SPL, f), data_only=True)
        ws = wb[wb.sheetnames[0]]
        for r in ws.iter_rows(min_row=4, values_only=True):
            if r and r[0]:
                want.add(clean(r[0]))

    print("fetching 01-10-2026 opening prices for %d symbols ..." % len(want))
    px, missing = fetch_opens(want)
    print("    got %d, missing %d%s" % (len(px), len(missing),
                                        (": " + ", ".join(missing)) if missing else ""))

    # ---- 2. the three equity books ----
    for f in HEDGE:
        path = os.path.join(SPL, f)
        wb = openpyxl.load_workbook(path)
        ws = wb["Port_" + SIG_NEW]
        hr, ci = table(ws)
        changed = unpriced = 0
        held = {}
        for r in range(hr + 1, ws.max_row + 1):
            raw = ws.cell(r, ci["Stock"]).value
            if not raw:
                continue
            s = clean(raw)
            w = num(ws.cell(r, ci["SIM Weight"]).value)
            prev = int(num(ws.cell(r, ci["Prev_Qty"]).value))
            if w <= 0:
                # an exit row: quantity stays 0, it is the prior holding that sells
                held[s] = (0, prev, 0.0)
                continue
            # No October bar from the source: the engine's own rule is to fall
            # back to the prior close (som_hedge.py line 517), so do exactly that
            # rather than leaving a quantity sized off nothing in particular.
            o = px[s][0] if s in px else sep_close(s)
            if not o or o <= 0:
                unpriced += 1
                held[s] = (int(num(ws.cell(r, ci["Qty"]).value)), prev, 0.0)
                continue
            q = max(1, int(np.floor(CAP * w / o)))
            old = int(num(ws.cell(r, ci["Qty"]).value))
            if q != old:
                changed += 1
            d = q - prev
            ws.cell(r, ci["Qty"], q)
            ws.cell(r, ci["Delta"], d)
            ws.cell(r, ci["Action"],
                    "HOLD" if d == 0 else ("BUY %d" % d if d > 0 else "SELL %d" % abs(d)))
            held[s] = (q, prev, w)

        # the live October row: gap P&L on the outgoing book plus the move since
        # the open, net of turnover cost -- the engine's own formula, now that a
        # real October open exists.
        wsm = wb["Detailed_Monthly_Summary"]
        head = [str(c.value).strip() if c.value is not None else "" for c in wsm[1]]
        mi = {h: i + 1 for i, h in enumerate(head)}
        row = None
        for r in range(2, wsm.max_row + 1):
            if str(wsm.cell(r, mi["Month"]).value).strip() == SIG_NEW:
                row = r
                break
        if row:
            # The outgoing (September) book, for the gap leg.
            sep = {}
            wsp = wb["Port_2026-08"]
            hr2, ci2 = table(wsp)
            for r in range(hr2 + 1, wsp.max_row + 1):
                raw = wsp.cell(r, ci2["Stock"]).value
                if not raw:
                    continue
                q = int(num(wsp.cell(r, ci2["Qty"]).value))
                if q > 0:
                    sep[clean(raw)] = q
            pnl = txn = 0.0
            for s in set(held) | set(sep):
                q = held.get(s, (0, 0, 0.0))[0]
                prev = sep.get(s, 0)
                pc = sep_close(s)
                o, last = px.get(s, (pc, pc))
                if not pc or not o:
                    continue
                pnl += (o - pc) * prev + (last - o) * q
                txn += abs(q - prev) * o
            live = (pnl - txn * COST_PER_TRADE) / CAP
            stored = num(wsm.cell(row, mi["Base"]).value)
            # Reported, NOT written. The stored live-month value is left as the
            # turnover cost the engine itself writes for a freshly formed book
            # ("Future Selection Mode"), because nothing refreshes that cell
            # until the month-end re-run -- a 09:18 snapshot would sit there as
            # October's return for the whole month. app.js renderHeatmap draws
            # the last monthly_detail row as a "still forming" marker rather
            # than a coloured return, and the live panel recomputes MTD from
            # live prices on every rebuild, so the stored number is not what
            # anyone reads anyway.
            print("  %-28s %3d qty re-sized, %d unpriced | stored %+.3f%% "
                  "(turnover cost) | marked-to-market now %+.3f%%"
                  % (f, changed, unpriced, stored * 100, live * 100))
        if not DRY:
            wb.save(path)

    # ---- 3. High Quality and ML current books ----
    for f in CURRENT_BOOKS:
        path = os.path.join(SPL, f)
        wb = openpyxl.load_workbook(path)
        ws = wb[wb.sheetnames[0]]
        changed = unpriced = 0
        r = 4
        while r <= ws.max_row and ws.cell(r, 1).value:
            s = clean(ws.cell(r, 1).value)
            w = num(ws.cell(r, 6).value)
            prev = int(num(ws.cell(r, 5).value))
            if w <= 0:                              # exit row
                if s in px:
                    ws.cell(r, 7, px[s][0])
                r += 1
                continue
            o = px[s][0] if s in px else sep_close(s)
            if not o or o <= 0:
                unpriced += 1
                r += 1
                continue
            q = max(1, int(np.floor(CAP * w / o)))
            if q != int(num(ws.cell(r, 4).value)):
                changed += 1
            d = q - prev
            ws.cell(r, 4, q)
            ws.cell(r, 3, abs(d))
            ws.cell(r, 2, "HOLD" if d == 0 else ("BUY" if d > 0 else "SELL"))
            ws.cell(r, 7, o)
            r += 1
        print("  %-44s %3d qty re-sized, %d unpriced" % (f, changed, unpriced))
        if not DRY:
            wb.save(path)

    print("\n%s (%s)" % ("DRY RUN - nothing written" if DRY
                         else "rewrote " + SPL,
                         datetime.now().strftime("%Y-%m-%d %H:%M:%S IST")))


if __name__ == "__main__":
    main()
