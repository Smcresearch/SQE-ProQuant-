"""
build_october_hq.py
-------------------
Splice October onto the PUBLISHED High Quality book, same principle as
build_october_books.py: history is immutable, only the tail is rewritten.

Three things are wrong with simply shipping the October run:

  1. Its history moves. The live workbook was produced from a 282-file
     hq_quarterly_universe; scratch/_backfill_hq_universe.py then added 20
     price files so the 16 names eligible for 2026-09 could be priced at all.
     Those names were also eligible in earlier quarters, so re-running picks
     them up retroactively -- PDSL appears in 2023-12 and 311 cells of history
     change. Correct going forward, but not something to restate.

  2. The live September is only half a month. Its Bench Return % is -3.39%
     while September actually closed -5.88%; the row was written before the
     month finished. Today it is hidden as the "live" row, but on 01-10 the
     calendar filter promotes it into CAGR/Sharpe/drawdown as if complete.

  3. Four names the screen dropped (SKYGOLD, PGIL, MANORAMA, ANUP) just vanish
     from the October selection with no sell instruction.

So: history and the realised September come from a re-run on the ORIGINAL 282
files (scratch/hq_282, verified to reproduce the live history to 0.0028pp), and
the October selection comes from the full 302-file run, which is the only one
that can see AGIIL/BUILDPRO/SIGMAADV -- three names October actually holds. The
two runs' capital bases agree to 0.11%, so the selection is used as-is; only
Prev Qty is re-pointed at the real September book, and the four dropped names
get explicit exit rows.

    python build_october_hq.py            # -> scratch/spliced/
    python build_october_hq.py --install  # also copy over the live workbooks
"""
import os
import re
import shutil
import sys

import openpyxl
import pandas as pd

MAIN = os.path.dirname(os.path.abspath(__file__))
HIST = os.path.join(MAIN, "scratch", "hq_282")     # 282 files, END_MONTH=2026-09
FULL = os.path.join(MAIN, "scratch", "oct_run")    # 302 files, END_MONTH=2026-09
OUT = os.path.join(MAIN, "scratch", "spliced")

SUMMARY = "SOM_HQ_Quarterly_v2_Summary.xlsx"
WORKBOOK = "SOM_HQ_Quarterly_v2.xlsx"
CURRENT = "SOM_HQ_Quarterly_v2_Current.xlsx"

SIG_REALIZE = "2026-08"   # held September
SIG_NEW = "2026-09"       # held October
MONTH_RE = re.compile(r"^\d{4}-\d{2}$")


def hdr_row(ws):
    return next(r for r in range(1, 15)
                if ws.cell(r, 1).value and "Portfolio" in str(ws.cell(r, 1).value))


def cols(ws, hr):
    out = {}
    for c in range(1, ws.max_column + 1):
        v = ws.cell(hr, c).value
        if v:
            out[str(v).replace("\n", " ").strip()] = c
    return out


def month_rows(ws, hr):
    return {str(ws.cell(r, 1).value).strip(): r
            for r in range(hr + 1, ws.max_row + 1)
            if ws.cell(r, 1).value and MONTH_RE.match(str(ws.cell(r, 1).value).strip())}


def num(v):
    try:
        f = float(v)
        return 0.0 if f != f else f
    except (TypeError, ValueError):
        return 0.0


def selection(path):
    """Selection_ sheet -> {symbol: {tgt, prev, w, px}} plus the sheet title."""
    wb = openpyxl.load_workbook(path, data_only=True)
    ws = wb[wb.sheetnames[0]]
    out = {}
    for r in ws.iter_rows(min_row=4, values_only=True):
        if not r or not r[0]:
            continue
        s = str(r[0]).replace("_1d_max", "").strip()
        out[s] = {"tgt": int(num(r[3])), "prev": int(num(r[4])),
                  "w": num(r[5]), "px": num(r[6])}
    return ws.title, out, ws["A1"].value


def main():
    install = "--install" in sys.argv
    os.makedirs(OUT, exist_ok=True)
    for d in (HIST, FULL):
        if not os.path.isdir(d):
            sys.exit("[error] missing run directory " + d)

    # ================= 1. Summary workbook =================
    src = os.path.join(MAIN, SUMMARY)
    dst = os.path.join(OUT, SUMMARY)
    shutil.copy2(src, dst)

    wbh = openpyxl.load_workbook(os.path.join(HIST, SUMMARY), data_only=True)
    wsh = wbh["Summary 5Y"]
    hrh = hdr_row(wsh)
    ch = cols(wsh, hrh)
    mrh = month_rows(wsh, hrh)

    wbf = openpyxl.load_workbook(os.path.join(FULL, SUMMARY), data_only=True)
    wsf = wbf["Summary 5Y"]
    hrf = hdr_row(wsf)
    cf = cols(wsf, hrf)
    mrf = month_rows(wsf, hrf)

    wb = openpyxl.load_workbook(dst)
    ws = wb["Summary 5Y"]
    hr = hdr_row(ws)
    ci = cols(ws, hr)
    mr = month_rows(ws, hr)

    if SIG_REALIZE not in mr:
        sys.exit("[error] live Summary has no row for " + SIG_REALIZE)
    prev_sig = sorted(m for m in mr if m < SIG_REALIZE)[-1]

    # -- realise September, chaining cumulative off the LIVE previous row so
    #    history stays exactly as published rather than as the re-run has it.
    r_live, r_hist = mr[SIG_REALIZE], mrh[SIG_REALIZE]
    old_ret = num(ws.cell(r_live, ci["Port Return %"]).value)
    old_bench = num(ws.cell(r_live, ci["Bench Return %"]).value)
    carry = {k: num(ws.cell(mr[prev_sig], ci[k]).value)
             for k in ("Port Cumul %", "Bench Cumul %", "Compounded Capital (Rs.)")}

    for k in ("Stocks", "Added Stocks", "Removed Stocks", "RF Rate", "Ex-ante Beta",
              "Invest (Rs.)", "Exit Value (Rs.)", "Net PnL (Rs.)", "Port Return %",
              "Bench Return %", "Ex-ante Vol", "Ex-ante Sharpe",
              "Added Symbols", "Removed Symbols"):
        if k in ci and k in ch:
            ws.cell(r_live, ci[k], wsh.cell(r_hist, ch[k]).value)
    ret = num(ws.cell(r_live, ci["Port Return %"]).value)
    ben = num(ws.cell(r_live, ci["Bench Return %"]).value)
    pc = (1 + carry["Port Cumul %"]) * (1 + ret) - 1
    bc = (1 + carry["Bench Cumul %"]) * (1 + ben) - 1
    ws.cell(r_live, ci["Alpha %"], ret - ben)
    if "Beat Bench?" in ci:
        ws.cell(r_live, ci["Beat Bench?"], "YES" if ret > ben else "NO")
    ws.cell(r_live, ci["Port Cumul %"], pc)
    ws.cell(r_live, ci["Bench Cumul %"], bc)
    cap = carry["Compounded Capital (Rs.)"] * (1 + ret)
    ws.cell(r_live, ci["Compounded Capital (Rs.)"], cap)
    if "Growth of Rs.100" in ci:
        ws.cell(r_live, ci["Growth of Rs.100"], 100 * (1 + pc))
    print("=== High Quality ===")
    print("    Sept (sig %s) realised: %+.3f%% -> %+.3f%%   bench %+.3f%% -> %+.3f%%"
          % (SIG_REALIZE, old_ret * 100, ret * 100, old_bench * 100, ben * 100))

    # -- the October selection, and the churn against the REAL September book
    sep_name, sep, _ = selection(os.path.join(MAIN, CURRENT))
    oct_name, octb, oct_title = selection(os.path.join(FULL, CURRENT))
    added = sorted(s for s in octb if s not in sep)
    exits = sorted(s for s in sep if s not in octb)

    # -- insert the October row ahead of the trailing blank + TOTAL rows
    ins = max(mr.values()) + 1
    ws.insert_rows(ins)
    r_new = mrf[SIG_NEW]
    for k, c in ci.items():
        if k in cf:
            ws.cell(ins, c, wsf.cell(r_new, cf[k]).value)
    ws.cell(ins, ci["Portfolio Month"], SIG_NEW)
    if "Trade Month" in ci:
        ws.cell(ins, ci["Trade Month"],
                (pd.Period(SIG_NEW, "M") + 1).strftime("%Y-%m"))
    ws.cell(ins, ci["Stocks"], len(octb))
    ws.cell(ins, ci["Added Stocks"], len(added))
    ws.cell(ins, ci["Removed Stocks"], len(exits))
    for k in ("Port Return %", "Bench Return %", "Alpha %"):
        ws.cell(ins, ci[k], 0)
    ws.cell(ins, ci["Port Cumul %"], pc)
    ws.cell(ins, ci["Bench Cumul %"], bc)
    ws.cell(ins, ci["Compounded Capital (Rs.)"], cap)
    if "Growth of Rs.100" in ci:
        ws.cell(ins, ci["Growth of Rs.100"], 100 * (1 + pc))
    if "Added Symbols" in ci:
        ws.cell(ins, ci["Added Symbols"], ", ".join(added))
    if "Removed Symbols" in ci:
        ws.cell(ins, ci["Removed Symbols"], ", ".join(exits))
    if "Beat Bench?" in ci:
        ws.cell(ins, ci["Beat Bench?"], "")
    wb.save(dst)
    print("    Oct row (sig %s / trade %s): %d holdings, %d added, %d exited"
          % (SIG_NEW, (pd.Period(SIG_NEW, "M") + 1).strftime("%Y-%m"),
             len(octb), len(added), len(exits)))
    print("      added : " + ", ".join(added))
    print("      exited: " + ", ".join(exits))
    print("    -> " + dst)

    # ================= 2. PM sheets =================
    wsrc = os.path.join(MAIN, WORKBOOK)
    wdst = os.path.join(OUT, WORKBOOK)
    shutil.copy2(wsrc, wdst)
    wbw = openpyxl.load_workbook(wdst)
    for tag, srcbook, sheet in (("realised Sept", os.path.join(HIST, WORKBOOK),
                                 "PM_" + SIG_REALIZE),
                                ("October", os.path.join(FULL, WORKBOOK),
                                 "PM_" + SIG_NEW)):
        sb = openpyxl.load_workbook(srcbook, data_only=True)
        if sheet not in sb.sheetnames:
            print("    [warn] %s missing from %s" % (sheet, srcbook))
            continue
        if sheet in wbw.sheetnames:
            del wbw[sheet]
        s = sb[sheet]
        t = wbw.create_sheet(sheet)
        for row in s.iter_rows():
            for cell in row:
                if cell.value is not None:
                    t.cell(cell.row, cell.column, cell.value)
        print("    %s: %s <- %s" % (tag, sheet, os.path.basename(
            os.path.dirname(srcbook))))
    # Keep the month sheets in chronological order so "last PM_" stays meaningful.
    pm = sorted(s for s in wbw.sheetnames if re.fullmatch(r"PM_\d{4}-\d{2}", s))
    for s in pm:
        wbw.move_sheet(s, offset=len(wbw.sheetnames) - wbw.sheetnames.index(s) - 1)
    wbw.save(wdst)
    print("    -> " + wdst)

    # ================= 3. Current selection =================
    cdst = os.path.join(OUT, CURRENT)
    shutil.copy2(os.path.join(FULL, CURRENT), cdst)
    wbc = openpyxl.load_workbook(cdst)
    wsc = wbc[wbc.sheetnames[0]]
    fixed = 0
    r = 4
    while r <= wsc.max_row:
        s = wsc.cell(r, 1).value
        if not s:
            break
        sym = str(s).replace("_1d_max", "").strip()
        tgt = int(num(wsc.cell(r, 4).value))
        real_prev = sep.get(sym, {}).get("tgt", 0)
        if int(num(wsc.cell(r, 5).value)) != real_prev:
            fixed += 1
        wsc.cell(r, 5, real_prev)
        d = tgt - real_prev
        wsc.cell(r, 3, abs(d))
        wsc.cell(r, 2, "HOLD" if d == 0 else ("BUY" if d > 0 else "SELL"))
        r += 1
    for sym in exits:
        p = sep[sym]["tgt"]
        wsc.cell(r, 1, sym)
        wsc.cell(r, 2, "EXIT SELL")
        wsc.cell(r, 3, p)
        wsc.cell(r, 4, 0)
        wsc.cell(r, 5, p)
        wsc.cell(r, 6, 0.0)
        wsc.cell(r, 7, sep[sym]["px"])
        r += 1
    wbc.save(cdst)
    print("    Current: %s, Prev Qty corrected on %d row(s), %d exit row(s) added"
          % (oct_name, fixed, len(exits)))
    print("    -> " + cdst)

    if install:
        for f in (SUMMARY, WORKBOOK, CURRENT):
            shutil.copy2(os.path.join(OUT, f), os.path.join(MAIN, f))
            print("    installed -> " + os.path.join(MAIN, f))
    else:
        print("\nnot installed; rerun with --install")


if __name__ == "__main__":
    main()
